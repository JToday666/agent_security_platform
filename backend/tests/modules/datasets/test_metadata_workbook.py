from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.modules.datasets.metadata_registry import load_metadata_bundle
from app.modules.datasets.metadata_workbook import export_metadata_workbook, sync_metadata_from_workbook
from tests.helpers.dataset_bundle import write_minimal_registry


def test_json_to_xlsx_to_json_round_trip_updates_structured_fields(tmp_path: Path) -> None:
    registry_root = tmp_path / "dataset_metadata"
    write_minimal_registry(registry_root)

    bundle = load_metadata_bundle(registry_root)
    workbook_path = registry_root / "workbook" / "dataset_registry.xlsx"
    export_metadata_workbook(bundle, workbook_path)

    workbook = load_workbook(workbook_path)
    subtype_sheet = workbook["risk_subtypes"]
    headers = [cell.value for cell in subtype_sheet[1]]
    assert headers == ["code", "category_code", "name", "sort_order", "is_active"]
    assert subtype_sheet.cell(row=2, column=4).value == 1
    assert subtype_sheet.cell(row=2, column=5).value is True

    display_sheet = workbook["risk_subtype_display_meta"]
    display_sheet.cell(row=2, column=2).value = "新的短描述"
    display_sheet.cell(row=2, column=3).value = "新的长描述"

    resources_sheet = workbook["risk_subtype_resources"]
    resources_sheet.append(["A1_identity_leakage", 2, "补充文档", "https://example.com/docs/A1-extra", "docs"])

    media_sheet = workbook["risk_subtype_media"]
    media_sheet.append(
        [
            "A1_identity_leakage",
            2,
            "A1-video",
            "video",
            "演示视频",
            "展示风险场景。",
            "https://example.com/media/A1.mp4",
            "https://example.com/media/A1-cover.png",
            2,
        ]
    )
    workbook.save(workbook_path)

    sync_metadata_from_workbook(workbook_path, registry_root)
    updated_bundle = load_metadata_bundle(registry_root)
    display_meta = updated_bundle.display_meta_by_code["A1_identity_leakage"]

    assert display_meta.short_description == "新的短描述"
    assert display_meta.full_description == "新的长描述"
    assert len(display_meta.resources) == 2
    assert display_meta.resources[1]["label"] == "补充文档"
    assert len(display_meta.media) == 2
    assert display_meta.media[1]["media_id"] == "A1-video"

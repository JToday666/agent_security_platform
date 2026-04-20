"""数据集元数据的 JSON <-> XLSX 转换逻辑。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.modules.datasets.importer import ImportValidationError
from app.modules.datasets.metadata_registry import (
    AssetTypeRecord,
    AttackDeliveryTypeRecord,
    DatasetSourceRecord,
    DisplayMetaRecord,
    MetadataBundle,
    RiskCategoryRecord,
    RiskSubtypeRecord,
    load_metadata_bundle,
    write_metadata_bundle,
)


def export_metadata_workbook(bundle: MetadataBundle, xlsx_path: Path) -> None:
    """把 JSON 元数据导出为工作簿。"""
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["sheet", "说明"])
    readme.append(["dataset_sources", "来源字典表"])
    readme.append(["attack_delivery_types", "攻击投递方式字典表"])
    readme.append(["asset_types", "资产类型字典表"])
    readme.append(["risk_categories", "风险大类"])
    readme.append(["risk_subtypes", "风险子类"])
    readme.append(["risk_subtype_display_meta", "短描述与长描述"])
    readme.append(["risk_subtype_highlights", "亮点列表"])
    readme.append(["risk_subtype_scenarios", "场景列表"])
    readme.append(["risk_subtype_resources", "资源列表"])
    readme.append(["risk_subtype_media", "媒体列表"])

    _append_sheet(
        workbook,
        "dataset_sources",
        ["code", "name", "description", "is_active"],
        ([item.code, item.name, item.description, item.is_active] for item in sorted(bundle.dataset_sources, key=lambda item: item.code)),
    )
    _append_sheet(
        workbook,
        "attack_delivery_types",
        ["code", "name", "description", "is_active"],
        (
            [item.code, item.name, item.description, item.is_active]
            for item in sorted(bundle.attack_delivery_types, key=lambda item: item.code)
        ),
    )
    _append_sheet(
        workbook,
        "asset_types",
        ["code", "name", "description", "is_active"],
        ([item.code, item.name, item.description, item.is_active] for item in sorted(bundle.asset_types, key=lambda item: item.code)),
    )
    _append_sheet(
        workbook,
        "risk_categories",
        ["code", "name", "meaning", "description", "sort_order", "is_active"],
        (
            [item.code, item.name, item.meaning, item.description, item.sort_order, item.is_active]
            for item in sorted(bundle.risk_categories, key=lambda item: ((item.sort_order or 999), item.code))
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtypes",
        ["code", "category_code", "name", "sort_order", "is_active"],
        (
            [item.code, item.category_code, item.name, item.sort_order, item.is_active]
            for item in sorted(bundle.risk_subtypes, key=lambda item: ((item.sort_order or 999), item.code))
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtype_display_meta",
        ["subtype_code", "short_description", "full_description"],
        (
            [item.subtype_code, item.short_description, item.full_description]
            for item in _sorted_display_meta(bundle)
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtype_highlights",
        ["subtype_code", "seq_no", "text"],
        (
            [item.subtype_code, index, text]
            for item in _sorted_display_meta(bundle)
            for index, text in enumerate(item.highlights, start=1)
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtype_scenarios",
        ["subtype_code", "seq_no", "text"],
        (
            [item.subtype_code, index, text]
            for item in _sorted_display_meta(bundle)
            for index, text in enumerate(item.scenarios, start=1)
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtype_resources",
        ["subtype_code", "seq_no", "label", "url", "type"],
        (
            [item.subtype_code, index, resource.get("label"), resource.get("url"), resource.get("type")]
            for item in _sorted_display_meta(bundle)
            for index, resource in enumerate(item.resources, start=1)
        ),
    )
    _append_sheet(
        workbook,
        "risk_subtype_media",
        ["subtype_code", "seq_no", "media_id", "type", "title", "description", "url", "cover_url", "sort"],
        (
            [
                item.subtype_code,
                index,
                media.get("media_id"),
                media.get("type"),
                media.get("title"),
                media.get("description"),
                media.get("url"),
                media.get("cover_url"),
                media.get("sort"),
            ]
            for item in _sorted_display_meta(bundle)
            for index, media in enumerate(item.media, start=1)
        ),
    )

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def sync_metadata_from_workbook(xlsx_path: Path, registry_root: Path) -> MetadataBundle:
    """从工作簿回写 JSON 元数据。"""
    workbook = load_workbook(xlsx_path)
    bundle = load_metadata_bundle(registry_root)

    bundle.dataset_sources = _read_simple_sheet(
        workbook["dataset_sources"],
        lambda row: DatasetSourceRecord(
            code=_require_text(row[0], "dataset_sources.code"),
            name=_require_text(row[1], "dataset_sources.name"),
            description=_optional_text(row[2]),
            is_active=_coerce_bool(row[3], "dataset_sources.is_active"),
        ),
    )
    bundle.attack_delivery_types = _read_simple_sheet(
        workbook["attack_delivery_types"],
        lambda row: AttackDeliveryTypeRecord(
            code=_require_text(row[0], "attack_delivery_types.code"),
            name=_require_text(row[1], "attack_delivery_types.name"),
            description=_optional_text(row[2]),
            is_active=_coerce_bool(row[3], "attack_delivery_types.is_active"),
        ),
    )
    bundle.asset_types = _read_simple_sheet(
        workbook["asset_types"],
        lambda row: AssetTypeRecord(
            code=_require_text(row[0], "asset_types.code"),
            name=_require_text(row[1], "asset_types.name"),
            description=_optional_text(row[2]),
            is_active=_coerce_bool(row[3], "asset_types.is_active"),
        ),
    )
    bundle.risk_categories = _read_simple_sheet(
        workbook["risk_categories"],
        lambda row: RiskCategoryRecord(
            code=_require_text(row[0], "risk_categories.code"),
            name=_require_text(row[1], "risk_categories.name"),
            meaning=_optional_text(row[2]),
            description=_optional_text(row[3]),
            sort_order=_coerce_optional_int(row[4], "risk_categories.sort_order"),
            is_active=_coerce_bool(row[5], "risk_categories.is_active"),
        ),
    )
    bundle.risk_subtypes = _read_simple_sheet(
        workbook["risk_subtypes"],
        lambda row: RiskSubtypeRecord(
            code=_require_text(row[0], "risk_subtypes.code"),
            category_code=_require_text(row[1], "risk_subtypes.category_code"),
            name=_require_text(row[2], "risk_subtypes.name"),
            sort_order=_coerce_optional_int(row[3], "risk_subtypes.sort_order"),
            is_active=_coerce_bool(row[4], "risk_subtypes.is_active"),
        ),
    )

    existing_meta = dict(bundle.display_meta_by_code)
    bundle.display_meta_by_code = {}
    for row in workbook["risk_subtype_display_meta"].iter_rows(min_row=2, values_only=True):
        if _is_blank_row(row):
            continue
        subtype_code = _require_text(row[0], "risk_subtype_display_meta.subtype_code")
        existing = existing_meta.get(subtype_code, DisplayMetaRecord(subtype_code=subtype_code))
        bundle.display_meta_by_code[subtype_code] = DisplayMetaRecord(
            subtype_code=subtype_code,
            short_description=_optional_text(row[1]),
            full_description=_optional_text(row[2]),
            highlights=[],
            scenarios=[],
            resources=[],
            media=[],
        )
        if existing and subtype_code not in bundle.display_meta_by_code:
            bundle.display_meta_by_code[subtype_code] = existing

    for subtype_code in existing_meta:
        bundle.display_meta_by_code.setdefault(subtype_code, existing_meta[subtype_code])
        bundle.display_meta_by_code[subtype_code].highlights = []
        bundle.display_meta_by_code[subtype_code].scenarios = []
        bundle.display_meta_by_code[subtype_code].resources = []
        bundle.display_meta_by_code[subtype_code].media = []

    for subtype_code in bundle.display_meta_by_code:
        bundle.display_meta_by_code[subtype_code].highlights = []
        bundle.display_meta_by_code[subtype_code].scenarios = []
        bundle.display_meta_by_code[subtype_code].resources = []
        bundle.display_meta_by_code[subtype_code].media = []

    _fill_string_list_sheet(workbook["risk_subtype_highlights"], bundle.display_meta_by_code, "highlights")
    _fill_string_list_sheet(workbook["risk_subtype_scenarios"], bundle.display_meta_by_code, "scenarios")
    _fill_resources_sheet(workbook["risk_subtype_resources"], bundle.display_meta_by_code)
    _fill_media_sheet(workbook["risk_subtype_media"], bundle.display_meta_by_code)

    write_metadata_bundle(registry_root, bundle)
    return bundle


def _append_sheet(workbook: Workbook, title: str, headers: list[str], rows) -> None:
    """向工作簿追加一个带表头的数据工作表。"""
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))


def _sorted_display_meta(bundle: MetadataBundle) -> list[DisplayMetaRecord]:
    """按 subtype code 排序展示元数据记录。"""
    return [bundle.display_meta_by_code[key] for key in sorted(bundle.display_meta_by_code)]


def _read_simple_sheet(sheet, factory):
    """按行读取简单结构工作表并映射为记录对象。"""
    items = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_blank_row(row):
            continue
        items.append(factory(row))
    return items


def _fill_string_list_sheet(sheet, display_meta_by_code: dict[str, DisplayMetaRecord], field_name: str) -> None:
    """把字符串列表工作表回填到展示元数据对象。"""
    grouped: dict[str, list[tuple[int, str]]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_blank_row(row):
            continue
        subtype_code = _require_text(row[0], f"{sheet.title}.subtype_code")
        seq_no = _coerce_required_int(row[1], f"{sheet.title}.seq_no")
        text = _require_text(row[2], f"{sheet.title}.text")
        _require_display_meta(display_meta_by_code, subtype_code, sheet.title)
        grouped.setdefault(subtype_code, []).append((seq_no, text))

    for subtype_code, items in grouped.items():
        items.sort(key=lambda item: item[0])
        setattr(display_meta_by_code[subtype_code], field_name, [text for _, text in items])


def _fill_resources_sheet(sheet, display_meta_by_code: dict[str, DisplayMetaRecord]) -> None:
    """把资源列表工作表回填到展示元数据对象。"""
    grouped: dict[str, list[tuple[int, dict[str, object]]]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_blank_row(row):
            continue
        subtype_code = _require_text(row[0], "risk_subtype_resources.subtype_code")
        seq_no = _coerce_required_int(row[1], "risk_subtype_resources.seq_no")
        _require_display_meta(display_meta_by_code, subtype_code, sheet.title)
        grouped.setdefault(subtype_code, []).append(
            (
                seq_no,
                {
                    "label": _require_text(row[2], "risk_subtype_resources.label"),
                    "url": _require_text(row[3], "risk_subtype_resources.url"),
                    "type": _require_text(row[4], "risk_subtype_resources.type"),
                },
            )
        )
    for subtype_code, items in grouped.items():
        items.sort(key=lambda item: item[0])
        display_meta_by_code[subtype_code].resources = [item for _, item in items]


def _fill_media_sheet(sheet, display_meta_by_code: dict[str, DisplayMetaRecord]) -> None:
    """把媒体列表工作表回填到展示元数据对象。"""
    grouped: dict[str, list[tuple[int, dict[str, object]]]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_blank_row(row):
            continue
        subtype_code = _require_text(row[0], "risk_subtype_media.subtype_code")
        seq_no = _coerce_required_int(row[1], "risk_subtype_media.seq_no")
        _require_display_meta(display_meta_by_code, subtype_code, sheet.title)
        grouped.setdefault(subtype_code, []).append(
            (
                seq_no,
                {
                    "media_id": _require_text(row[2], "risk_subtype_media.media_id"),
                    "type": _require_text(row[3], "risk_subtype_media.type"),
                    "title": _require_text(row[4], "risk_subtype_media.title"),
                    "description": _optional_text(row[5]),
                    "url": _require_text(row[6], "risk_subtype_media.url"),
                    "cover_url": _optional_text(row[7]),
                    "sort": _coerce_optional_int(row[8], "risk_subtype_media.sort"),
                },
            )
        )
    for subtype_code, items in grouped.items():
        items.sort(key=lambda item: item[0])
        display_meta_by_code[subtype_code].media = [item for _, item in items]


def _require_display_meta(display_meta_by_code: dict[str, DisplayMetaRecord], subtype_code: str, sheet_name: str) -> None:
    """校验工作表引用的 subtype 已在展示元数据表中定义。"""
    if subtype_code not in display_meta_by_code:
        raise ImportValidationError(f"{sheet_name}: subtype_code={subtype_code} 未在 risk_subtype_display_meta 中定义")


def _require_text(value, field_name: str) -> str:
    """读取必填文本单元格并校验非空。"""
    if value is None:
        raise ImportValidationError(f"{field_name} 不能为空")
    if not isinstance(value, str):
        value = str(value)
    normalized = value.strip()
    if not normalized:
        raise ImportValidationError(f"{field_name} 不能为空")
    return normalized


def _optional_text(value) -> str | None:
    """读取可选文本单元格并去除空白。"""
    if value is None:
        return None
    if not isinstance(value, str):
        value = str(value)
    normalized = value.strip()
    return normalized or None


def _coerce_bool(value, field_name: str) -> bool:
    """校验工作表中的布尔值字段。"""
    if isinstance(value, bool):
        return value
    raise ImportValidationError(f"{field_name} 必须是布尔值")


def _coerce_optional_int(value, field_name: str) -> int | None:
    """读取可选整数单元格。"""
    if value is None or value == "":
        return None
    if isinstance(value, int):
        return value
    raise ImportValidationError(f"{field_name} 必须是整数")


def _coerce_required_int(value, field_name: str) -> int:
    """读取必填整数单元格。"""
    if isinstance(value, int):
        return value
    raise ImportValidationError(f"{field_name} 必须是整数")


def _is_blank_row(row) -> bool:
    """判断一行单元格是否全部为空。"""
    return all(cell is None or cell == "" for cell in row)
